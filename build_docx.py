"""
build_docx.py — 학습교재 Markdown → 정식 Word(.docx) 변환기
================================================
· 목차(TOC 필드), 제목 페이지, Word 제목 스타일, 진짜 표
· ```mermaid 다이어그램은 matplotlib로 이미지 렌더링해 삽입(순서도/시퀀스/트리)
· 코드/트리 블록은 고정폭(Consolas), 인용은 콜아웃 박스, 링크는 하이퍼링크
실행:  .\.venv\Scripts\python.exe build_docx.py
"""
from __future__ import annotations

import os
import re
import shutil

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Polygon, FancyArrowPatch

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT

SRC = "유니버셜키오스크_학습교재.md"
OUT = "유니버셜키오스크_학습교재.docx"
ASSET = "_docx_assets"

KFONT = "맑은 고딕"          # 한글 본문
CODEFONT = "Consolas"
plt.rcParams["font.family"] = "Malgun Gothic"
plt.rcParams["axes.unicode_minus"] = False

EMOJI = re.compile(
    "[\U0001F000-\U0001FAFF\U00002600-\U000027BF\U0001F1E6-\U0001F1FF\u2190-\u21FF\u2B00-\u2BFF\uFE0F\u200d]")


def strip_emoji(s: str) -> str:
    return EMOJI.sub("", s).strip()


# ──────────────────────────────────────────────
# LaTeX(KaTeX) → 읽기 좋은 유니코드
# ──────────────────────────────────────────────
_SUB = {"0": "\u2080", "1": "\u2081", "2": "\u2082", "i": "\u1D62", "j": "\u2C7C", "n": "\u2099"}


def latex_to_unicode(s: str) -> str:
    s = s.strip()
    s = s.replace("\\left", "").replace("\\right", "").replace("\\,", " ")
    s = re.sub(r"\\mathbf\{(.*?)\}", r"\1", s)
    s = re.sub(r"\\hat\{(.*?)\}", "\\1\u0302", s)
    s = re.sub(r"\\bar\{(.*?)\}", "\\1\u0304", s)
    s = re.sub(r"\\widehat\{(.*?)\}", "\\1\u0302", s)
    s = re.sub(r"\\sqrt\{(.*?)\}", "\u221A(\\1)", s)
    s = re.sub(r"\\frac\{(.*?)\}\{(.*?)\}", r"(\1)/(\2)", s)
    s = re.sub(r"\\sum_\{.*?\}\^\{.*?\}", "\u03A3", s)
    s = re.sub(r"\\sum", "\u03A3", s)
    s = s.replace("\\lVert", "\u2016").replace("\\rVert", "\u2016")
    s = s.replace("\\cdot", "\u00B7").replace("\\theta", "\u03B8")
    s = s.replace("\\times", "\u00D7").replace("\\approx", "\u2248")
    s = re.sub(r"\^2", "\u00B2", s)
    s = re.sub(r"_\{([a-zA-Z0-9]+)\}", lambda m: "(" + m.group(1) + ")", s)
    s = re.sub(r"_([0-9ijn])", lambda m: _SUB.get(m.group(1), "_" + m.group(1)), s)
    s = s.replace("{", "").replace("}", "")
    return s.strip()


# ──────────────────────────────────────────────
# docx 저수준 헬퍼
# ──────────────────────────────────────────────
def set_cell_bg(cell, hexcolor: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def set_para_shading(p, hexcolor: str):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    pPr.append(shd)


def set_left_border(p, hexcolor="2D6CDF", size="18"):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), size)
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), hexcolor)
    pbdr.append(left)
    pPr.append(pbdr)


def add_hyperlink(paragraph, url: str, text: str):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "0563C1"); rPr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rPr.append(u)
    rf = OxmlElement("w:rFonts"); rf.set(qn("w:eastAsia"), KFONT); rPr.append(rf)
    new_run.append(rPr)
    t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|https?://[^\s)]+|\$[^$]+\$)")


def add_inline(paragraph, text: str, bold=False, italic=False, color=None, font=None, size=None):
    def style(run):
        run.bold = bold
        run.italic = italic
        if color:
            run.font.color.rgb = color
        run.font.name = font or KFONT
        run._element.rPr.rFonts.set(qn("w:eastAsia"), font or KFONT)
        if size:
            run.font.size = Pt(size)

    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            r = paragraph.add_run(piece[2:-2]); r.bold = True; r.italic = italic
            r.font.name = font or KFONT; r._element.rPr.rFonts.set(qn("w:eastAsia"), font or KFONT)
            if color: r.font.color.rgb = color
            if size: r.font.size = Pt(size)
        elif piece.startswith("`") and piece.endswith("`"):
            r = paragraph.add_run(piece[1:-1]); r.font.name = CODEFONT
            r._element.rPr.rFonts.set(qn("w:eastAsia"), CODEFONT)
            r.font.color.rgb = RGBColor(0xC0, 0x37, 0x22)
            if size: r.font.size = Pt(size)
        elif piece.startswith("[") and "](" in piece:
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", piece)
            add_hyperlink(paragraph, m.group(2), m.group(1))
        elif piece.startswith("http"):
            add_hyperlink(paragraph, piece, piece)
        elif piece.startswith("$") and piece.endswith("$"):
            r = paragraph.add_run(latex_to_unicode(piece[1:-1])); r.italic = True
            r.font.name = "Cambria Math"; r._element.rPr.rFonts.set(qn("w:eastAsia"), KFONT)
            if size: r.font.size = Pt(size)
        else:
            r = paragraph.add_run(piece)
            style(r)


def add_toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "[여기에 목차가 생성됩니다 — Word에서 Ctrl+A → F9 로 업데이트]"
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    for el in (f1, instr, f2, t, f3):
        run._r.append(el)


def enable_update_fields(doc):
    settings = doc.settings.element
    upd = OxmlElement("w:updateFields"); upd.set(qn("w:val"), "true")
    settings.append(upd)


def set_base_fonts(doc):
    for name in ("Normal",):
        st = doc.styles[name]
        st.font.name = KFONT
        st.font.size = Pt(10.5)
        rpr = st.element.get_or_add_rPr()
        rf = rpr.get_or_add_rFonts()
        rf.set(qn("w:eastAsia"), KFONT)
        rf.set(qn("w:ascii"), KFONT)
        rf.set(qn("w:hAnsi"), KFONT)
    for hn in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        try:
            st = doc.styles[hn]
            rpr = st.element.get_or_add_rPr()
            rf = rpr.get_or_add_rFonts()
            rf.set(qn("w:eastAsia"), KFONT)
        except KeyError:
            pass


# ──────────────────────────────────────────────
# 다이어그램 렌더링 (matplotlib)
# ──────────────────────────────────────────────
def _wrap(label, width=11):
    out = []
    for part in label.split("\\n"):
        part = strip_emoji(part)
        line = ""
        for word in part.split(" "):
            while len(word) > width:            # 긴 토큰 강제 분할
                if line:
                    out.append(line); line = ""
                out.append(word[:width]); word = word[width:]
            if len(line) + len(word) + 1 > width and line:
                out.append(line); line = word
            else:
                line = (line + " " + word).strip()
        if line:
            out.append(line)
    return [x for x in out if x != ""] or [" "]


def parse_flow(src):
    lines = [l.rstrip() for l in src.splitlines() if l.strip()]
    direction = "TD"
    nodes, shape, edges, order = {}, {}, [], []

    def add_node(nid, label=None, shp=None):
        if nid not in nodes:
            nodes[nid] = label if label is not None else nid
            order.append(nid)
        elif label is not None:
            nodes[nid] = label
        if shp:
            shape[nid] = shp

    ndef = re.compile(r"([A-Za-z0-9_가-힣]+)(\[\(.*?\)\]|\(\[.*?\]\)|\{\{.*?\}\}|\[.*?\]|\{.*?\})")

    def lab_shape(d):
        if d.startswith("{{"): return d[2:-2], "rect"
        if d.startswith("{"): return d[1:-1], "diamond"
        if d.startswith("[("): return d[2:-2], "round"
        if d.startswith("(["): return d[2:-2], "round"
        if d.startswith("["): return d[1:-1], "rect"
        return d, "rect"

    for ln in lines:
        low = ln.strip().lower()
        if low.startswith("flowchart") or low.startswith("graph"):
            for tok in ln.split():
                if tok.upper() in ("TD", "TB", "LR", "RL", "BT"):
                    direction = "LR" if tok.upper() in ("LR", "RL") else "TD"
            continue
        if low.startswith("statediagram"):
            direction = "TD"; continue
        for m in ndef.finditer(ln):
            lbl, shp = lab_shape(m.group(2))
            add_node(m.group(1), lbl, shp)
        stripped = ndef.sub(lambda m: " " + m.group(1) + " ", ln)
        stripped = stripped.replace("[*]", " _START_ ")
        if "_START_" in stripped:
            add_node("_START_", "\u25CF", "round")
        if "-->" not in stripped and "<-->" not in stripped:
            continue
        conn = stripped.replace("<-->", "-->")
        elabel = ""
        m = re.search(r"-->\|(.*?)\|", conn)
        if m:
            elabel = m.group(1); conn = re.sub(r"-->\|.*?\|", " --> ", conn)
        m = re.search(r"--\s+([^>|]+?)\s+-->", conn)
        if m:
            elabel = elabel or m.group(1).strip(); conn = re.sub(r"--\s+[^>|]+?\s+-->", " --> ", conn)
        if "-->" not in conn:
            continue
        left, right = conn.split("-->", 1)
        if ":" in right:
            right, sl = right.split(":", 1); elabel = elabel or sl.strip()
        srcs = [x.strip() for x in left.split("&") if x.strip()]
        dsts = [x.strip() for x in right.split("&") if x.strip()]
        for s in srcs:
            for d in dsts:
                s2 = s.split()[-1] if s.split() else ""
                d2 = d.split()[0] if d.split() else ""
                if s2 and d2:
                    add_node(s2); add_node(d2)
                    edges.append((s2, d2, strip_emoji(elabel)))
    return direction, nodes, shape, edges, order


def _forward_layers(nodes, edges):
    from collections import defaultdict
    adj = defaultdict(list)
    for s, d, _ in edges:
        adj[s].append(d)
    color = {n: 0 for n in nodes}          # 0 white, 1 gray, 2 black
    back = set()
    import sys
    sys.setrecursionlimit(10000)

    def dfs(u):
        color[u] = 1
        for v in adj[u]:
            if color[v] == 1:
                back.add((u, v))
            elif color[v] == 0:
                dfs(v)
        color[u] = 2

    for nid in list(nodes):
        if color[nid] == 0:
            dfs(nid)
    fedges = [(s, d) for s, d, _ in edges if (s, d) not in back]
    layer = {n: 0 for n in nodes}
    for _ in range(len(nodes) + 1):
        changed = False
        for s, d in fedges:
            if layer[d] < layer[s] + 1:
                layer[d] = layer[s] + 1; changed = True
        if not changed:
            break
    return layer, back


def _order_layers(nodes, edges, layer):
    from collections import defaultdict
    layers = defaultdict(list)
    for n in nodes:
        layers[layer[n]].append(n)
    maxL = max(layer.values()) if layer else 0
    pred = defaultdict(list); succ = defaultdict(list)
    for s, d, _ in edges:
        if layer.get(d) == layer.get(s, -9) + 1:
            succ[s].append(d); pred[d].append(s)
    pos = {}
    for L in range(maxL + 1):
        for i, n in enumerate(layers[L]):
            pos[n] = i
    for _ in range(4):
        for L in range(1, maxL + 1):
            layers[L].sort(key=lambda n: (sum(pos[p] for p in pred[n]) / len(pred[n])) if pred[n] else pos[n])
            for i, n in enumerate(layers[L]):
                pos[n] = i
        for L in range(maxL - 1, -1, -1):
            layers[L].sort(key=lambda n: (sum(pos[c] for c in succ[n]) / len(succ[n])) if succ[n] else pos[n])
            for i, n in enumerate(layers[L]):
                pos[n] = i
    return layers, maxL


def draw_flow(src, path):
    direction, nodes, shape, edges, order = parse_flow(src)
    if not nodes:
        return False
    layer, back = _forward_layers(nodes, edges)
    layers, maxL = _order_layers(nodes, edges, layer)
    wrapped = {n: _wrap(nodes[n], 11) for n in nodes}

    def nw(n):
        return max(1.5, 0.16 * max(len(x) for x in wrapped[n]) + 0.5)

    def nh(n):
        return max(0.6, 0.36 * len(wrapped[n]) + 0.36)

    maxw = max(nw(n) for n in nodes)
    maxh = max(nh(n) for n in nodes)
    pos = {}
    if direction == "TD":
        slot = maxw + 0.55; gap = maxh + 1.0
        for L in range(maxL + 1):
            row = layers[L]
            for i, n in enumerate(row):
                pos[n] = ((i - (len(row) - 1) / 2) * slot, -L * gap)
    else:
        slot = maxh + 0.5; gap = maxw + 1.5
        for L in range(maxL + 1):
            col = layers[L]
            for i, n in enumerate(col):
                pos[n] = (L * gap, -(i - (len(col) - 1) / 2) * slot)
    xs = [p[0] for p in pos.values()]; ys = [p[1] for p in pos.values()]
    minx, maxx = min(xs), max(xs); miny, maxy = min(ys), max(ys)
    fig_w = min(6.6, max(3.2, (maxx - minx) + maxw + 0.8))
    fig_h = min(9.4, max(2.2, (maxy - miny) + maxh + 0.8))
    fig, ax = plt.subplots(figsize=(fig_w, fig_h))
    ax.set_xlim(minx - maxw / 2 - 0.4, maxx + maxw / 2 + 0.4)
    ax.set_ylim(miny - maxh / 2 - 0.4, maxy + maxh / 2 + 0.4)
    ax.axis("off")
    for s, d, el in edges:
        if s not in pos or d not in pos:
            continue
        x1, y1 = pos[s]; x2, y2 = pos[d]
        is_back = (s, d) in back
        if direction == "TD" and not is_back:
            start = (x1, y1 - nh(s) / 2); end = (x2, y2 + nh(d) / 2)
            rad = 0.0 if abs(x1 - x2) < 0.05 else 0.06
        elif direction == "LR" and not is_back:
            start = (x1 + nw(s) / 2, y1); end = (x2 - nw(d) / 2, y2)
            rad = 0.0 if abs(y1 - y2) < 0.05 else 0.06
        else:
            start = (x1 + nw(s) / 2, y1); end = (x2 + nw(d) / 2, y2); rad = 0.45
        ax.add_patch(FancyArrowPatch(start, end, arrowstyle="-|>", mutation_scale=12,
                                     lw=1.25, color="#5B6B7F",
                                     connectionstyle=f"arc3,rad={rad}"))
        if el:
            if is_back:
                lx = max(start[0], end[0]) + 0.5; ly = (start[1] + end[1]) / 2
            else:
                lx = start[0] * 0.5 + end[0] * 0.5
                ly = start[1] * 0.5 + end[1] * 0.5
                if direction == "LR":
                    ly += 0.24 if end[1] >= start[1] else -0.24
            ax.text(lx, ly, el, fontsize=7.0, color="#2D6CDF", ha="center", va="center",
                    bbox=dict(boxstyle="round,pad=0.12", fc="white", ec="none", alpha=0.92))
    for n in order:
        if n not in pos:
            continue
        x, y = pos[n]; w = nw(n); h = nh(n)
        txt = "\n".join(wrapped[n])
        shp = shape.get(n, "rect")
        if n == "_START_":
            ax.add_patch(plt.Circle((x, y), 0.15, fc="#2D6CDF", ec="none"))
            continue
        if shp == "diamond":
            ax.add_patch(Polygon([(x, y + h / 1.4), (x + w / 1.6, y), (x, y - h / 1.4), (x - w / 1.6, y)],
                                 closed=True, fc="#FFF4D6", ec="#E6A700", lw=1.4))
        else:
            fc = "#E8F0FE" if shp == "rect" else "#E7F7EC"
            ax.add_patch(FancyBboxPatch((x - w / 2, y - h / 2), w, h,
                        boxstyle="round,pad=0.02,rounding_size=0.1", fc=fc, ec="#2D6CDF", lw=1.25))
        ax.text(x, y, txt, ha="center", va="center", fontsize=7.8, color="#12203A")
    fig.savefig(path, dpi=170, bbox_inches="tight")
    plt.close(fig)
    return True


def parse_sequence(src):
    parts, names, msgs = [], {}, []
    for ln in src.splitlines():
        ln = ln.strip()
        m = re.match(r"participant\s+(\S+)\s+as\s+(.+)", ln)
        if m:
            parts.append(m.group(1)); names[m.group(1)] = strip_emoji(m.group(2)); continue
        m = re.match(r"(\S+?)\s*(->>|-->>|->|-->)\s*(\S+?)\s*:\s*(.+)", ln)
        if m:
            a, arrow, b, txt = m.groups()
            for p in (a, b):
                if p not in parts:
                    parts.append(p); names.setdefault(p, p)
            msgs.append((a, b, strip_emoji(txt), arrow in ("-->>", "-->")))
    return parts, names, msgs


def draw_sequence(src, path):
    parts, names, msgs = parse_sequence(src)
    if not parts or not msgs:
        return False
    xg = 2.6
    x = {p: i * xg for i, p in enumerate(parts)}
    top = 0.0
    dy = 0.85
    bottom = -(len(msgs) + 1) * dy
    W = max(4.0, (len(parts) - 1) * xg + 2.2)
    H = max(2.5, (len(msgs) + 2) * dy)
    fig, ax = plt.subplots(figsize=(min(7.0, W), min(9.2, H)))
    ax.set_xlim(-1.3, (len(parts) - 1) * xg + 1.3)
    ax.set_ylim(bottom - 0.6, 0.9)
    ax.axis("off")
    for p in parts:
        ax.add_patch(FancyBboxPatch((x[p] - 1.0, top - 0.05), 2.0, 0.55,
                    boxstyle="round,pad=0.02,rounding_size=0.1", fc="#2D6CDF", ec="none"))
        ax.text(x[p], top + 0.22, "\n".join(_wrap(names[p], 12)), ha="center", va="center",
                fontsize=8.5, color="white")
        ax.plot([x[p], x[p]], [top - 0.05, bottom], color="#B7C2D2", lw=1.0, ls=(0, (4, 3)))
    for i, (a, b, txt, dashed) in enumerate(msgs):
        y = -(i + 1) * dy
        ax.annotate("", xy=(x[b], y), xytext=(x[a], y),
                    arrowprops=dict(arrowstyle="-|>", lw=1.3, color="#5B6B7F",
                                    linestyle="--" if dashed else "-"))
        midx = (x[a] + x[b]) / 2
        ax.text(midx, y + 0.14, txt, ha="center", va="bottom", fontsize=7.4, color="#12203A",
                bbox=dict(boxstyle="round,pad=0.15", fc="white", ec="none", alpha=0.9))
    fig.tight_layout(pad=0.2)
    fig.savefig(path, dpi=170, bbox_inches="tight")
    plt.close(fig)
    return True


def render_mermaid(src, idx):
    path = os.path.join(ASSET, f"diagram_{idx}.png")
    kind = src.strip().splitlines()[0].strip().lower()
    try:
        if kind.startswith("sequencediagram"):
            ok = draw_sequence(src, path)
        else:
            ok = draw_flow(src, path)
        return path if ok else None
    except Exception as e:
        print("  [diagram fail]", e)
        return None


# ──────────────────────────────────────────────
# 블록 렌더링 헬퍼
# ──────────────────────────────────────────────
def add_code_block(doc, code_lines):
    p = doc.add_paragraph()
    set_para_shading(p, "F4F5F7")
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    first = True
    for ln in code_lines:
        r = p.add_run(("" if first else "\n") + ln.replace("\t", "    "))
        r.font.name = CODEFONT
        r._element.rPr.rFonts.set(qn("w:eastAsia"), CODEFONT)
        r.font.size = Pt(8.8)
        first = False


def add_quote_block(doc, qlines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    set_para_shading(p, "EEF3FB")
    set_left_border(p)
    for i, ln in enumerate(qlines):
        if i:
            p.add_run().add_break()
        add_inline(p, ln, size=9.5)


def add_table(doc, rows):
    header = rows[0]
    body = rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, cell_text in enumerate(header):
        c = t.rows[0].cells[j]
        set_cell_bg(c, "2D6CDF")
        p = c.paragraphs[0]
        add_inline(p, cell_text, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=9.5)
    for row in body:
        cells = t.add_row().cells
        for j in range(len(header)):
            txt = row[j] if j < len(row) else ""
            p = cells[j].paragraphs[0]
            add_inline(p, txt, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def split_table_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


# ──────────────────────────────────────────────
# 메인 변환
# ──────────────────────────────────────────────
def main():
    if os.path.isdir(ASSET):
        shutil.rmtree(ASSET)
    os.makedirs(ASSET, exist_ok=True)

    with open(SRC, "r", encoding="utf-8") as f:
        lines = f.read().split("\n")

    # front matter(제목 페이지) / 본문 분리
    body_start = next((i for i, l in enumerate(lines) if l.startswith("## ")), 0)
    front = lines[:body_start]
    body = lines[body_start:]

    doc = Document()
    set_base_fonts(doc)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.9)
    sec.top_margin = sec.bottom_margin = Inches(0.9)

    # ── 제목 페이지 ──
    title_txt = "유니버셜 키오스크 — 결선 대비 학습 교재"
    for l in front:
        if l.startswith("# "):
            title_txt = strip_emoji(l[2:])
    doc.add_paragraph().paragraph_format.space_after = Pt(60)
    tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = tp.add_run(title_txt); tr.bold = True; tr.font.size = Pt(26)
    tr.font.color.rgb = RGBColor(0x1B, 0x23, 0x30)
    tr.font.name = KFONT
    tr._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), KFONT)
    for l in front:
        if l.strip().startswith(">"):
            q = l.strip().lstrip(">").strip()
            if q:
                pp = doc.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(pp, q, size=11)
    doc.add_paragraph().paragraph_format.space_after = Pt(30)
    sp = doc.add_paragraph(); sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(sp, "**기술 학습 · 구술 심사 대비 문서**", size=12)
    doc.add_page_break()

    # ── 목차 ──
    h = doc.add_heading("목차", level=1)
    add_toc(doc)
    doc.add_page_break()

    # ── 본문 파싱 ──
    i = 0
    n = len(body)
    dia_idx = 0
    while i < n:
        line = body[i]
        s = line.strip()

        if s.startswith("```"):
            lang = s[3:].strip().lower()
            block = []
            i += 1
            while i < n and not body[i].strip().startswith("```"):
                block.append(body[i]); i += 1
            i += 1  # closing fence
            if lang == "mermaid":
                dia_idx += 1
                img = render_mermaid("\n".join(block), dia_idx)
                if img and os.path.exists(img):
                    pic = doc.add_paragraph(); pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    from PIL import Image
                    w_px = Image.open(img).size[0]
                    width_in = min(6.4, w_px / 170.0)
                    pic.add_run().add_picture(img, width=Inches(width_in))
                else:
                    add_code_block(doc, block)
            else:
                add_code_block(doc, block)
            continue

        if s.startswith("|") and (i + 1 < n) and set(body[i + 1].strip()) <= set("|-: "):
            rows = []
            while i < n and body[i].strip().startswith("|"):
                if not set(body[i].strip()) <= set("|-: "):
                    rows.append(split_table_row(body[i]))
                i += 1
            add_table(doc, rows)
            continue

        if s.startswith(">"):
            qlines = []
            while i < n and body[i].strip().startswith(">"):
                qlines.append(body[i].strip().lstrip(">").strip())
                i += 1
            add_quote_block(doc, qlines)
            continue

        if s.startswith("#### "):
            doc.add_heading(strip_emoji(s[5:]), level=3); i += 1; continue
        if s.startswith("### "):
            doc.add_heading(strip_emoji(s[4:]), level=2); i += 1; continue
        if s.startswith("## "):
            doc.add_heading(strip_emoji(s[3:]), level=1); i += 1; continue
        if s.startswith("# "):
            i += 1; continue

        if s == "---" or s == "":
            i += 1; continue

        # 리스트
        m = re.match(r"^(\s*)([-*])\s+(.*)", line)
        if m:
            indent = len(m.group(1)); text = m.group(3)
            style = "List Bullet 2" if indent >= 2 else "List Bullet"
            text = re.sub(r"^\[[ xX]\]\s*", lambda mm: "\u2610 ", text)
            try:
                p = doc.add_paragraph(style=style)
            except KeyError:
                p = doc.add_paragraph(style="List Bullet")
            add_inline(p, text)
            i += 1; continue
        m = re.match(r"^(\s*)\d+\.\s+(.*)", line)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, m.group(2))
            i += 1; continue

        # 일반 문단
        p = doc.add_paragraph()
        add_inline(p, s)
        i += 1

    enable_update_fields(doc)
    doc.save(OUT)
    print("SAVED:", OUT, "| diagrams:", dia_idx)


if __name__ == "__main__":
    main()
