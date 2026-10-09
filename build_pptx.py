"""
build_pptx.py — 초등학생용 발표 PPT + 퀴즈 100 자동 생성
================================================
· 발표용.pptx : 기술을 쉬운 비유로, 슬라이드마다 matplotlib 설명그림
· 퀴즈100.pptx : 10주제 × 10문제, 문제 다음 슬라이드에 정답
그림은 matplotlib로 직접 그림(오프라인, 친근한 도식). 한글=맑은 고딕.
실행:  .\.venv\Scripts\python.exe build_pptx.py
"""
from __future__ import annotations
import os, shutil, math

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import (FancyBboxPatch, FancyArrowPatch, Circle, Arc,
                                Rectangle, Polygon, Wedge)

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

plt.rcParams["font.family"] = "Malgun Gothic"
plt.rcParams["axes.unicode_minus"] = False

ASSET = "_pptx_assets"

# 팔레트
BG = "#EEF4FF"
PRIMARY = "#2D6CDF"
ACCENT = "#FF8A3D"
GREEN = "#2BB673"
YELLOW = "#FFC93C"
PINK = "#FF6B9D"
INK = "#1B2330"
GRAY = "#8A9BAE"

RGB = lambda h: RGBColor(int(h[1:3], 16), int(h[3:5], 16), int(h[5:7], 16))


# ══════════════════════════════════════════════
# matplotlib 그림 도구
# ══════════════════════════════════════════════
def fig_ax(bg=BG):
    fig, ax = plt.subplots(figsize=(12.8, 7.2), dpi=150)
    ax.set_xlim(0, 16); ax.set_ylim(0, 9); ax.axis("off")
    ax.add_patch(Rectangle((0, 0), 16, 9, fc=bg, ec="none", zorder=-10))
    return fig, ax


def save(fig, name):
    path = os.path.join(ASSET, name + ".png")
    fig.savefig(path, dpi=150, bbox_inches="tight", pad_inches=0.1)
    plt.close(fig)
    return path


def rbox(ax, x, y, w, h, fc="white", ec=PRIMARY, lw=2.5, rad=0.28, z=1):
    ax.add_patch(FancyBboxPatch((x, y), w, h,
                 boxstyle=f"round,pad=0.02,rounding_size={rad}", fc=fc, ec=ec, lw=lw, zorder=z))


def T(ax, x, y, s, fs=17, color=INK, bold=True, ha="center", va="center", z=4, rot=0):
    ax.text(x, y, s, fontsize=fs, color=color, ha=ha, va=va,
            fontweight="bold" if bold else "normal", zorder=z, rotation=rot)


def arrow(ax, p1, p2, color=INK, lw=3, rad=0.0, z=2):
    ax.add_patch(FancyArrowPatch(p1, p2, arrowstyle="-|>", mutation_scale=24,
                 lw=lw, color=color, connectionstyle=f"arc3,rad={rad}", zorder=z))


def wrap(s, width=14):
    out, line = [], ""
    for w in s.split(" "):
        if len(line) + len(w) + 1 > width and line:
            out.append(line); line = w
        else:
            line = (line + " " + w).strip()
    if line:
        out.append(line)
    return "\n".join(out)


# ── 마스코트 로봇 ──
def mascot(ax, cx, cy, s=1.0, say=None, mood="happy"):
    rbox(ax, cx - 1.15 * s, cy - 0.1 * s, 2.3 * s, 1.7 * s, fc="#EAF1FF", ec=PRIMARY, lw=3, rad=0.45, z=3)
    ax.add_patch(Circle((cx - 0.42 * s, cy + 0.68 * s), 0.19 * s, fc=PRIMARY, zorder=4))
    ax.add_patch(Circle((cx + 0.42 * s, cy + 0.68 * s), 0.19 * s, fc=PRIMARY, zorder=4))
    ax.add_patch(Circle((cx - 0.36 * s, cy + 0.74 * s), 0.06 * s, fc="white", zorder=5))
    ax.add_patch(Circle((cx + 0.48 * s, cy + 0.74 * s), 0.06 * s, fc="white", zorder=5))
    if mood == "happy":
        ax.add_patch(Arc((cx, cy + 0.32 * s), 0.95 * s, 0.7 * s, theta1=200, theta2=340, lw=3, color=PRIMARY, zorder=4))
    else:
        ax.add_patch(Circle((cx, cy + 0.25 * s), 0.12 * s, fc="none", ec=PRIMARY, lw=3, zorder=4))
    ax.plot([cx, cx], [cy + 1.6 * s, cy + 2.1 * s], color=PRIMARY, lw=3, zorder=3)
    ax.add_patch(Circle((cx, cy + 2.2 * s), 0.14 * s, fc=ACCENT, zorder=4))
    rbox(ax, cx - 0.85 * s, cy - 1.5 * s, 1.7 * s, 1.3 * s, fc=PRIMARY, ec=PRIMARY, rad=0.3, z=2)
    ax.add_patch(Circle((cx, cy - 0.85 * s), 0.16 * s, fc="white", zorder=3))
    if say:
        bubble(ax, cx + 1.35 * s, cy + 0.9 * s, say)


def bubble(ax, x, y, text, fs=15):
    txt = wrap(text, 16)
    lines = txt.count("\n") + 1
    w = min(6.2, 0.28 * max(len(l) for l in txt.split("\n")) + 1.0)
    h = 0.55 * lines + 0.5
    rbox(ax, x, y - h / 2, w, h, fc="#FFFDF3", ec=ACCENT, lw=2.5, rad=0.3, z=5)
    ax.add_patch(Polygon([(x, y - 0.1), (x - 0.5, y - 0.5), (x + 0.05, y - 0.55)],
                closed=True, fc="#FFFDF3", ec=ACCENT, lw=2.5, zorder=5))
    T(ax, x + w / 2, y, txt, fs=fs, color="#8a5a12", z=6)


# ── 작은 심볼들 ──
def sym_face(ax, cx, cy, r=0.9, fc="#FFE7C2"):
    ax.add_patch(Circle((cx, cy), r, fc=fc, ec=INK, lw=2.5, zorder=3))
    ax.add_patch(Circle((cx - r * 0.35, cy + r * 0.2), r * 0.12, fc=INK, zorder=4))
    ax.add_patch(Circle((cx + r * 0.35, cy + r * 0.2), r * 0.12, fc=INK, zorder=4))
    ax.add_patch(Arc((cx, cy - r * 0.15), r * 0.9, r * 0.7, theta1=200, theta2=340, lw=2.5, color=INK, zorder=4))


def sym_camera(ax, cx, cy, s=1.0):
    rbox(ax, cx - 1.2 * s, cy - 0.8 * s, 2.4 * s, 1.6 * s, fc="#2A3340", ec="#2A3340", rad=0.2, z=3)
    ax.add_patch(Circle((cx, cy), 0.55 * s, fc="#7FC7FF", ec="white", lw=3, zorder=4))
    ax.add_patch(Circle((cx, cy), 0.24 * s, fc="#2A3340", zorder=5))
    rbox(ax, cx + 0.5 * s, cy + 0.55 * s, 0.6 * s, 0.4 * s, fc="#FF5B5B", ec="#FF5B5B", rad=0.1, z=4)


def sym_hand_point(ax, cx, cy, s=1.0):
    rbox(ax, cx - 0.35 * s, cy - 1.0 * s, 0.7 * s, 1.3 * s, fc="#FFD9A8", ec=INK, lw=2, rad=0.25, z=3)
    rbox(ax, cx - 0.18 * s, cy + 0.1 * s, 0.36 * s, 1.5 * s, fc="#FFD9A8", ec=INK, lw=2, rad=0.18, z=4)


def sym_number(ax, cx, cy, s="123", fc=GREEN):
    rbox(ax, cx - 0.9, cy - 0.5, 1.8, 1.0, fc=fc, ec=fc, rad=0.25, z=3)
    T(ax, cx, cy, s, fs=20, color="white", z=4)


def sym_lock(ax, cx, cy, s=1.0):
    rbox(ax, cx - 0.6 * s, cy - 0.6 * s, 1.2 * s, 1.0 * s, fc=YELLOW, ec=INK, lw=2.5, rad=0.15, z=4)
    ax.add_patch(Arc((cx, cy + 0.35 * s), 0.8 * s, 0.9 * s, theta1=0, theta2=180, lw=3.5, color=INK, zorder=3))
    ax.add_patch(Circle((cx, cy - 0.1 * s), 0.12 * s, fc=INK, zorder=5))


# ══════════════════════════════════════════════
# 발표 슬라이드용 그림들
# ══════════════════════════════════════════════
def sc_cover():
    fig, ax = fig_ax("#DCEBFF")
    for i, c in enumerate(["#FFD1E8", "#D7F3E3", "#FFE9C7"]):
        ax.add_patch(Circle((2 + i * 5.6, 7.6), 0.5, fc=c, ec="none", zorder=0))
    mascot(ax, 3.2, 4.2, 1.4, say="안녕! 나랑 같이 배워보자!")
    sym_camera(ax, 12.3, 5.6, 1.1)
    sym_face(ax, 12.3, 2.7, 1.0)
    T(ax, 9.2, 8.0, "유니버셜 키오스크", fs=30, color=PRIMARY)
    T(ax, 9.2, 7.1, "쉽게 배우는 우리 작품 기술 이야기", fs=18, color=INK)
    return save(fig, "cover")


def sc_problem():
    fig, ax = fig_ax()
    rbox(ax, 5.5, 2.2, 5.0, 5.2, fc="#2A3340", ec="#2A3340", rad=0.3)
    T(ax, 8.0, 6.9, "키오스크", fs=18, color="white")
    for i in range(3):
        rbox(ax, 6.0 + i * 0.05, 3.0 + i * 1.1, 4.0, 0.9, fc="#3C4a5c", ec="#556", rad=0.15)
    sym_face(ax, 2.6, 4.6, 1.0, fc="#FFE7C2")
    bubble(ax, 0.8, 6.4, "글씨가 작아 안 보여요…")
    sym_face(ax, 13.4, 4.6, 1.0, fc="#E7D6FF")
    T(ax, 13.4, 2.9, "?!", fs=26, color=ACCENT)
    T(ax, 8.0, 1.3, "어떤 분들은 키오스크가 어려워요", fs=20, color=INK)
    return save(fig, "problem")


def sc_solution():
    fig, ax = fig_ax()
    sym_camera(ax, 3.0, 5.2, 1.1)
    T(ax, 3.0, 3.6, "카메라 + AI", fs=17, color=PRIMARY)
    arrow(ax, (4.7, 5.0), (7.0, 5.0), color=ACCENT, lw=4)
    sym_face(ax, 8.4, 5.0, 1.1)
    T(ax, 8.4, 3.4, "나이대를 알아요", fs=16)
    arrow(ax, (9.9, 5.0), (12.0, 5.0), color=ACCENT, lw=4)
    rbox(ax, 12.3, 3.9, 3.0, 2.2, fc="#EAF7EF", ec=GREEN, lw=3)
    T(ax, 13.8, 5.0, "화면이\n스스로 변신!", fs=16, color=GREEN)
    mascot(ax, 2.4, 1.9, 0.7, say="누구나 쉽게!")
    T(ax, 9.5, 1.2, "우리가 카메라와 AI로 도와줘요", fs=19, color=INK)
    return save(fig, "solution")


def sc_bigpic():
    fig, ax = fig_ax()
    items = [("📷", "카메라", "#EAF1FF", PRIMARY), ("🧠", "두뇌(AI)", "#FFF1E6", ACCENT),
             ("🖥️", "화면", "#EAF7EF", GREEN)]
    xs = [2.6, 8.0, 13.4]
    for (ic, name, fc, ec), x in zip(items, xs):
        rbox(ax, x - 1.6, 3.6, 3.2, 2.6, fc=fc, ec=ec, lw=3)
        T(ax, x, 5.4, name, fs=20, color=ec)
        T(ax, x, 4.4, "○", fs=22, color=ec)
    arrow(ax, (4.4, 4.9), (6.2, 4.9), color=INK, lw=4)
    arrow(ax, (9.8, 4.9), (11.6, 4.9), color=INK, lw=4)
    sym_camera(ax, 2.6, 4.5, 0.6)
    sym_face(ax, 13.4, 4.6, 0.7)
    T(ax, 8.0, 1.6, "카메라로 보고 → 두뇌가 생각하고 → 화면이 바뀐다", fs=18, color=INK)
    mascot(ax, 8.0, 7.9, 0.5)
    return save(fig, "bigpic")


def sc_four():
    fig, ax = fig_ax()
    data = [("나이 맞춤", "화면 변신", PRIMARY),
            ("손으로 주문", "가리키고 집기", ACCENT), ("말로 주문", "여러 나라 말", PINK)]
    pos = [(8.0, 5.2), (4.3, 2.3), (11.7, 2.3)]
    for (t, s, c), (x, y) in zip(data, pos):
        rbox(ax, x - 3.0, y - 1.0, 6.0, 2.0, fc="white", ec=c, lw=3)
        ax.add_patch(Circle((x - 2.1, y), 0.55, fc=c, ec=c, zorder=3))
        T(ax, x - 2.1, y, "★", fs=20, color="white")
        T(ax, x + 0.4, y + 0.35, t, fs=18, color=c)
        T(ax, x + 0.4, y - 0.4, s, fs=14, color=INK)
    T(ax, 8.0, 8.0, "우리 작품의 3가지 마법", fs=22, color=INK)
    return save(fig, "four")


def sc_privacy():
    fig, ax = fig_ax()
    rbox(ax, 2.0, 3.0, 4.2, 3.4, fc="white", ec=GRAY, lw=3)
    sym_face(ax, 4.1, 4.7, 1.1)
    T(ax, 4.1, 2.4, "사진", fs=16, color=GRAY)
    ax.plot([2.3, 5.9], [3.3, 6.1], color="#FF5B5B", lw=6, zorder=6)
    ax.plot([5.9, 2.3], [3.3, 6.1], color="#FF5B5B", lw=6, zorder=6)
    arrow(ax, (6.6, 4.7), (8.4, 4.7), color=ACCENT, lw=4)
    rbox(ax, 8.8, 3.0, 5.4, 3.4, fc="#FFF9E6", ec=YELLOW, lw=3)
    sym_lock(ax, 10.2, 4.9, 1.0)
    T(ax, 12.4, 5.2, "얼굴 정보", fs=16, color="#8a6a12")
    T(ax, 12.4, 4.2, "저장 안 함", fs=18, color="#8a6a12")
    T(ax, 8.0, 1.4, "얼굴 사진은 저장하지 않아요 (비밀 지키기)", fs=18, color=INK)
    return save(fig, "privacy")


def sc_age():
    fig, ax = fig_ax()
    sym_camera(ax, 2.6, 5.2, 0.9)
    arrow(ax, (4.0, 5.0), (5.6, 5.0), color=ACCENT, lw=4)
    rbox(ax, 5.8, 4.0, 2.4, 2.0, fc="#FFF1E6", ec=ACCENT, lw=3)
    T(ax, 7.0, 5.0, "AI", fs=26, color=ACCENT)
    T(ax, 7.0, 3.6, "나이 맞히기", fs=13)
    outs = [("어린이", GREEN, 7.2), ("어른", PRIMARY, 5.0), ("어르신", PINK, 2.8)]
    for name, c, y in outs:
        arrow(ax, (8.3, 5.0), (10.2, y), color=GRAY, lw=2.5)
        rbox(ax, 10.4, y - 0.55, 3.4, 1.1, fc="white", ec=c, lw=3)
        T(ax, 12.1, y, name + " 화면", fs=16, color=c)
    T(ax, 8.0, 1.2, "AI가 나이대를 맞혀 화면을 골라줘요", fs=18, color=INK)
    return save(fig, "age")


def sc_variants():
    fig, ax = fig_ax()
    rbox(ax, 1.8, 2.6, 4.0, 4.4, fc="white", ec=PRIMARY, lw=3)
    T(ax, 3.8, 6.4, "보통", fs=16, color=PRIMARY)
    T(ax, 3.8, 4.6, "메뉴", fs=18, color=INK)
    rbox(ax, 6.4, 2.6, 4.0, 4.4, fc="#FFF9E6", ec=ACCENT, lw=3)
    T(ax, 8.4, 6.4, "어르신 모드", fs=16, color=ACCENT)
    T(ax, 8.4, 4.6, "메뉴", fs=30, color=INK)
    T(ax, 8.4, 3.3, "글씨 1.5배!", fs=14, color=ACCENT)
    rbox(ax, 11.0, 2.6, 4.0, 4.4, fc="#EAF7EF", ec=GREEN, lw=3)
    T(ax, 13.0, 6.4, "어린이 모드", fs=16, color=GREEN)
    T(ax, 13.0, 4.6, "메뉴", fs=24, color=INK)
    T(ax, 13.0, 3.3, "큰 버튼·쉬운 말", fs=13, color=GREEN)
    T(ax, 8.0, 1.2, "같은 메뉴도 사람에 맞게 변신해요", fs=18, color=INK)
    return save(fig, "variants")


def sc_aim():
    fig, ax = fig_ax()
    sym_hand_point(ax, 3.2, 4.6, 1.2)
    arrow(ax, (4.2, 5.2), (7.8, 6.0), color=ACCENT, lw=3, rad=-0.2)
    ax.add_patch(Circle((8.4, 6.1), 0.45, fc="none", ec=PRIMARY, lw=4, zorder=4))
    ax.add_patch(Circle((8.4, 6.1), 0.1, fc=PRIMARY, zorder=4))
    rbox(ax, 9.6, 3.2, 5.0, 3.6, fc="white", ec=GRAY, lw=2)
    for i in range(2):
        for j in range(2):
            rbox(ax, 9.9 + j * 2.4, 3.6 + i * 1.6, 2.1, 1.3, fc="#EAF1FF", ec=PRIMARY, lw=1.5, rad=0.15)
    mascot(ax, 3.0, 1.9, 0.6, say="손이 커서를 움직여!")
    T(ax, 9.0, 1.2, "에임: 손을 들면 커서가 따라와요", fs=18, color=INK)
    return save(fig, "aim")


def sc_pinch():
    fig, ax = fig_ax()
    ax.add_patch(Polygon([(3.0, 3.0), (2.4, 5.2), (3.2, 5.6)], closed=True, fc="#FFD9A8", ec=INK, lw=2, zorder=3))
    ax.add_patch(Polygon([(3.0, 3.0), (4.2, 4.8), (3.6, 5.4)], closed=True, fc="#FFD9A8", ec=INK, lw=2, zorder=3))
    ax.add_patch(Circle((3.3, 5.5), 0.18, fc=ACCENT, zorder=5))
    T(ax, 3.0, 2.2, "엄지 + 검지 = 집기!", fs=15, color=INK)
    arrow(ax, (5.0, 4.6), (7.2, 4.6), color=ACCENT, lw=4)
    rbox(ax, 7.6, 3.4, 6.4, 2.6, fc="#EAF7EF", ec=GREEN, lw=3)
    T(ax, 10.8, 5.1, "메뉴 담기!", fs=20, color=GREEN)
    T(ax, 10.8, 4.0, "주먹보다 안정적 (두 손끝이 잘 보여요)", fs=13, color=INK)
    T(ax, 8.0, 1.2, "핀치: 손끝을 집으면 담겨요", fs=18, color=INK)
    return save(fig, "pinch")


def sc_dwell():
    fig, ax = fig_ax()
    for i, (frac, c) in enumerate([(0.15, GRAY), (0.55, YELLOW), (1.0, GREEN)]):
        cx = 3.2 + i * 4.4
        rbox(ax, cx - 1.6, 3.2, 3.2, 2.6, fc="#EAF1FF", ec=PRIMARY, lw=2)
        T(ax, cx, 4.5, "메뉴", fs=16)
        ax.add_patch(Circle((cx, 4.5), 0.7, fc="none", ec="#DDD", lw=6, zorder=4))
        ax.add_patch(Wedge((cx, 4.5), 0.7, 90, 90 - 360 * frac, width=0.001, ec=c, lw=6, zorder=5))
        T(ax, cx, 2.6, f"{int(frac*100)}%", fs=15, color=c)
    T(ax, 8.0, 6.9, "가만히 기다리면 링이 차올라요", fs=16, color=INK)
    mascot(ax, 13.8, 6.6, 0.5)
    T(ax, 8.0, 1.2, "드웰: 잠깐 머물면 자동으로 선택!", fs=18, color=INK)
    return save(fig, "dwell")


def sc_filter():
    fig, ax = fig_ax()
    import numpy as np
    xs = np.linspace(0, 6, 200)
    shaky = 5.6 + 0.5 * np.sin(xs * 6) + 0.3 * np.sin(xs * 17)
    smooth = 5.6 + 0.5 * np.sin(xs * 6)
    ax.plot(2 + xs, shaky, color="#FF5B5B", lw=2)
    T(ax, 5.0, 7.2, "떨리는 손", fs=15, color="#FF5B5B")
    arrow(ax, (8.4, 5.6), (9.6, 5.6), color=ACCENT, lw=4)
    T(ax, 9.0, 6.3, "필터", fs=15, color=ACCENT)
    ax.plot(9.8 + xs * 0.7, 3.0 + (smooth - 5.6) * 1.2 + 0.0, color=GREEN, lw=4)
    T(ax, 12.0, 4.2, "부드러운 커서", fs=15, color=GREEN)
    mascot(ax, 3.0, 2.4, 0.6, say="흔들림을 매끈하게!")
    T(ax, 9.5, 1.2, "필터가 손 떨림을 부드럽게 만들어요", fs=18, color=INK)
    return save(fig, "filter")


def sc_vote():
    fig, ax = fig_ax()
    picks = ["어른", "어른", "어린이", "어른", "어른"]
    for i, p in enumerate(picks):
        c = GREEN if p == "어른" else PINK
        rbox(ax, 1.6 + i * 2.7, 4.4, 2.3, 1.6, fc="white", ec=c, lw=2.5)
        T(ax, 2.75 + i * 2.7, 5.2, p, fs=15, color=c)
        T(ax, 2.75 + i * 2.7, 6.4, "손", fs=13, color=GRAY)
    arrow(ax, (8.0, 4.1), (8.0, 3.2), color=INK, lw=4)
    rbox(ax, 5.6, 1.9, 4.8, 1.2, fc="#EAF7EF", ec=GREEN, lw=3)
    T(ax, 8.0, 2.5, "다수결 → '어른'!", fs=18, color=GREEN)
    T(ax, 8.0, 7.6, "여러 번 보고 가장 많은 걸로 정해요", fs=18, color=INK)
    return save(fig, "vote")


def sc_voice():
    fig, ax = fig_ax()
    ax.add_patch(Circle((2.8, 5.0), 0.7, fc="#2A3340", zorder=3))
    rbox(ax, 2.5, 3.2, 0.6, 1.6, fc="#2A3340", ec="#2A3340", rad=0.2, z=2)
    T(ax, 2.8, 6.4, "마이크", fs=14)
    arrow(ax, (3.9, 5.0), (5.5, 5.0), color=ACCENT, lw=4)
    rbox(ax, 5.7, 4.1, 3.6, 1.8, fc="white", ec=PRIMARY, lw=2.5)
    T(ax, 7.5, 5.0, '"불고기버거 하나"', fs=14, color=INK)
    arrow(ax, (9.5, 5.0), (11.1, 5.0), color=ACCENT, lw=4)
    rbox(ax, 11.3, 4.0, 3.4, 2.0, fc="#EAF7EF", ec=GREEN, lw=3)
    T(ax, 13.0, 5.0, "장바구니\n담기", fs=15, color=GREEN)
    mascot(ax, 3.0, 2.2, 0.6, say="말을 글자로!")
    T(ax, 9.0, 1.2, "말로 주문 → 글자 → 명령으로 바뀌어요", fs=18, color=INK)
    return save(fig, "voice")


def sc_langs():
    fig, ax = fig_ax()
    langs = [("한국어", "안녕하세요", PRIMARY), ("English", "Hello", GREEN), ("中文(중국어)", "니하오", PINK)]
    for i, (t, s, c) in enumerate(langs):
        x = 3.0 + i * 5.0
        rbox(ax, x - 1.9, 3.4, 3.8, 3.0, fc="white", ec=c, lw=3)
        T(ax, x, 5.6, t, fs=20, color=c)
        T(ax, x, 4.3, s, fs=16, color=INK)
    T(ax, 8.0, 7.6, "여러 나라 말을 알아듣고 화면도 바뀌어요", fs=18, color=INK)
    mascot(ax, 8.0, 1.9, 0.55)
    return save(fig, "langs")


def sc_brainbody():
    fig, ax = fig_ax()
    rbox(ax, 1.8, 2.4, 5.6, 4.6, fc="#FFF1E6", ec=ACCENT, lw=3)
    T(ax, 4.6, 6.4, "core = 두뇌", fs=18, color=ACCENT)
    for i, n in enumerate(["비전", "수학", "자연어", "저장"]):
        rbox(ax, 2.3 + (i % 2) * 2.6, 3.2 + (i // 2) * 1.3, 2.2, 1.0, fc="white", ec=ACCENT, lw=1.5, rad=0.15)
        T(ax, 3.4 + (i % 2) * 2.6, 3.7 + (i // 2) * 1.3, n, fs=13)
    rbox(ax, 8.6, 2.4, 5.6, 4.6, fc="#EAF1FF", ec=PRIMARY, lw=3)
    T(ax, 11.4, 6.4, "ui = 몸(화면)", fs=18, color=PRIMARY)
    for i, n in enumerate(["창", "카드", "커서", "테마"]):
        rbox(ax, 9.1 + (i % 2) * 2.6, 3.2 + (i // 2) * 1.3, 2.2, 1.0, fc="white", ec=PRIMARY, lw=1.5, rad=0.15)
        T(ax, 10.2 + (i % 2) * 2.6, 3.7 + (i // 2) * 1.3, n, fs=13)
    T(ax, 8.0, 1.2, "두뇌(core)와 몸(ui)을 나눠서 만들어요", fs=18, color=INK)
    return save(fig, "brainbody")


def sc_threads():
    fig, ax = fig_ax()
    T(ax, 8.0, 7.8, "요리사 여러 명 = 동시에 일하기(스레드)", fs=18, color=INK)
    jobs = [("카메라 보기", PRIMARY, 3.2), ("화면 그리기", GREEN, 8.0), ("소리 안내", PINK, 12.8)]
    for name, c, x in jobs:
        ax.add_patch(Circle((x, 5.4), 0.7, fc=c, ec=c, zorder=3))
        T(ax, x, 5.4, "○", fs=22, color="white")
        rbox(ax, x - 1.7, 3.4, 3.4, 1.4, fc="white", ec=c, lw=2.5)
        T(ax, x, 4.1, name, fs=15, color=c)
    T(ax, 8.0, 1.5, "한 명이 카메라를 봐도 화면은 안 멈춰요", fs=17, color=INK)
    return save(fig, "threads")


def sc_db():
    fig, ax = fig_ax()
    for i in range(3):
        rbox(ax, 4.6, 2.6 + i * 1.4, 6.8, 1.2, fc="white", ec=PRIMARY, lw=2.5, rad=0.12)
        ax.add_patch(Circle((5.4, 3.2 + i * 1.4), 0.15, fc=ACCENT, zorder=4))
    T(ax, 8.4, 6.4, "서랍장(SQLite)", fs=18, color=PRIMARY)
    T(ax, 8.4, 5.0, "주문: 불고기버거 세트", fs=14)
    T(ax, 8.4, 3.6, "요일·시간대 통계", fs=14)
    T(ax, 8.4, 2.2, "(얼굴은 저장 안 함!)", fs=13, color=GREEN)
    mascot(ax, 13.6, 6.4, 0.5)
    T(ax, 8.0, 1.0, "작은 서랍장에 '익명 주문 통계'만 기억해요", fs=18, color=INK)
    return save(fig, "db")


def sc_states():
    fig, ax = fig_ax()
    steps = [("환영", PRIMARY), ("메뉴", ACCENT), ("완료", GREEN)]
    xs = [3.0, 8.0, 13.0]
    for (name, c), x in zip(steps, xs):
        ax.add_patch(Circle((x, 5.0), 1.1, fc="white", ec=c, lw=4, zorder=3))
        T(ax, x, 5.0, name, fs=18, color=c)
    arrow(ax, (4.2, 5.0), (6.8, 5.0), color=INK, lw=4)
    arrow(ax, (9.2, 5.0), (11.8, 5.0), color=INK, lw=4)
    T(ax, 5.5, 5.6, "주문 시작", fs=13, color=GRAY)
    T(ax, 10.5, 5.6, "결제", fs=13, color=GRAY)
    arrow(ax, (13.0, 3.9), (3.0, 3.9), color=GRAY, lw=2.5, rad=0.25)
    T(ax, 8.0, 2.6, "잠시 후 처음으로", fs=13, color=GRAY)
    T(ax, 8.0, 7.6, "신호등처럼 화면 단계가 바뀌어요(상태머신)", fs=18, color=INK)
    return save(fig, "states")


def sc_demo():
    fig, ax = fig_ax()
    sym_camera(ax, 3.4, 5.2, 1.0)
    ax.plot([2.2, 4.6], [4.2, 6.2], color="#FF5B5B", lw=6, zorder=6)
    ax.plot([4.6, 2.2], [4.2, 6.2], color="#FF5B5B", lw=6, zorder=6)
    T(ax, 3.4, 3.4, "카메라 없음", fs=15, color="#FF5B5B")
    arrow(ax, (5.4, 5.0), (7.2, 5.0), color=ACCENT, lw=4)
    rbox(ax, 7.6, 3.4, 6.6, 3.2, fc="#EAF7EF", ec=GREEN, lw=3)
    T(ax, 10.9, 5.6, "데모 모드로 OK!", fs=18, color=GREEN)
    T(ax, 10.9, 4.4, "버튼·음성으로 전부 시연", fs=14, color=INK)
    mascot(ax, 3.2, 1.9, 0.6, say="언제나 동작!")
    T(ax, 9.5, 1.1, "카메라·인터넷 없어도 멈추지 않아요", fs=18, color=INK)
    return save(fig, "demo")


def sc_points():
    fig, ax = fig_ax()
    import numpy as np
    rng = np.random.default_rng(3)
    sym_face(ax, 4.0, 4.8, 1.9, fc="#FFF3E0")
    for _ in range(90):
        a = rng.uniform(0, 2 * math.pi); r = rng.uniform(0, 1.7)
        ax.add_patch(Circle((4.0 + r * math.cos(a), 4.8 + r * 0.95 * math.sin(a)), 0.045, fc=PRIMARY, zorder=6))
    T(ax, 4.0, 2.4, "얼굴 468점", fs=18, color=PRIMARY)
    hx, hy = 11.5, 4.8
    pts = [(hx, hy - 1.6)]
    for f in range(5):
        for k in range(4):
            pts.append((hx - 1.4 + f * 0.7, hy - 1.0 + k * 0.7))
    for p in pts:
        ax.add_patch(Circle(p, 0.09, fc=GREEN, zorder=5))
    for f in range(5):
        xs = [hx] + [hx - 1.4 + f * 0.7 for k in range(4)]
        ys = [hy - 1.6] + [hy - 1.0 + k * 0.7 for k in range(4)]
        ax.plot(xs, ys, color=GREEN, lw=1.5, zorder=4)
    T(ax, 11.5, 2.4, "손 21점", fs=18, color=GREEN)
    T(ax, 8.0, 1.1, "컴퓨터는 얼굴·손을 '점'으로 봐요", fs=18, color=INK)
    return save(fig, "points")


def sc_flow_all():
    fig, ax = fig_ax()
    steps = ["시작", "얼굴/손 보기", "화면 변신", "주문(손·말)", "결제", "완료"]
    cols = [GRAY, PRIMARY, ACCENT, GREEN, PINK, PRIMARY]
    for i, (s, c) in enumerate(zip(steps, cols)):
        x = 1.5 + (i % 3) * 5.0
        y = 5.6 if i < 3 else 2.8
        rbox(ax, x - 1.5, y - 0.8, 3.0, 1.6, fc="white", ec=c, lw=3)
        T(ax, x, y, s, fs=15, color=c)
    arrow(ax, (3.0, 5.6), (5.0, 5.6), color=INK, lw=3.5)
    arrow(ax, (8.0, 5.6), (10.0, 5.6), color=INK, lw=3.5)
    arrow(ax, (11.5, 4.8), (11.5, 3.6), color=INK, lw=3.5)
    arrow(ax, (10.0, 2.8), (8.0, 2.8), color=INK, lw=3.5)
    arrow(ax, (5.0, 2.8), (3.0, 2.8), color=INK, lw=3.5)
    T(ax, 8.0, 7.8, "우리 프로그램의 전체 순서", fs=20, color=INK)
    return save(fig, "flow_all")


def sc_recap():
    fig, ax = fig_ax()
    lines = ["1. 카메라는 나이대만 보고(사진은 안 저장)",
             "2. AI가 나이 맞혀 화면이 변신",
             "3. 손: 가리키기·집기·기다리기",
             "4. 흔들림은 필터·다수결로 잡기",
             "5. 말로도 주문, 카메라 없어도 OK"]
    for i, l in enumerate(lines):
        rbox(ax, 1.4, 6.3 - i * 1.15, 11.4, 0.95, fc="white", ec=PRIMARY, lw=2, rad=0.2)
        T(ax, 7.1, 6.77 - i * 1.15, l, fs=16, ha="center")
    mascot(ax, 14.2, 4.4, 0.7, say="이것만 알면 끝!")
    T(ax, 7.1, 8.0, "오늘의 핵심 5줄", fs=22, color=INK)
    return save(fig, "recap")


def sc_end():
    fig, ax = fig_ax("#DCEBFF")
    mascot(ax, 8.0, 4.4, 1.5, say="발표 잘 할 수 있어! 화이팅!")
    T(ax, 8.0, 8.0, "고맙습니다!", fs=30, color=PRIMARY)
    T(ax, 8.0, 1.3, "NO PROBLEM KIOSK", fs=18, color=INK)
    return save(fig, "end")


# ══════════════════════════════════════════════
# 퀴즈 주제 배너 그림
# ══════════════════════════════════════════════
def topic_banner(idx, title, color):
    fig, ax = fig_ax("#FFFFFF")
    rbox(ax, 0.4, 3.2, 15.2, 2.6, fc=color, ec=color, rad=0.4)
    ax.add_patch(Circle((2.4, 4.5), 1.0, fc="white", ec="white", zorder=3))
    T(ax, 2.4, 4.5, str(idx), fs=40, color=color)
    T(ax, 9.2, 4.5, title, fs=26, color="white")
    mascot(ax, 14.0, 4.5, 0.55)
    return save(fig, f"topic_{idx}")


def mascot_small():
    fig, ax = plt.subplots(figsize=(3, 3), dpi=150)
    ax.set_xlim(0, 6); ax.set_ylim(0, 6); ax.axis("off")
    mascot(ax, 3, 3, 1.1)
    path = os.path.join(ASSET, "mascot_small.png")
    fig.savefig(path, dpi=150, bbox_inches="tight", transparent=True)
    plt.close(fig)
    return path


# ══════════════════════════════════════════════
# PPTX 헬퍼
# ══════════════════════════════════════════════
def set_run(run, size, color=INK, bold=True, font="맑은 고딕"):
    run.font.size = Pt(size); run.font.bold = bold
    run.font.color.rgb = RGB(color) if isinstance(color, str) else color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    ea = rPr.makeelement(qn("a:ea"), {"typeface": font}); rPr.append(ea)


def bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGB(color)


def band(slide, text, color=PRIMARY, y=0.0, h=1.15, fs=30, tcolor="#FFFFFF"):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(y), Inches(13.333), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = RGB(color); shp.line.fill.background()
    tf = shp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text; set_run(r, fs, tcolor, True)
    return shp


def textbox(slide, x, y, w, h, items, size=20, color=INK, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, bold=True, spacing=1.15):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = spacing; p.space_after = Pt(6)
        if isinstance(it, dict):
            text, sz, col, bd = it["t"], it.get("s", size), it.get("c", color), it.get("b", bold)
        else:
            text, sz, col, bd = it, size, color, bold
        r = p.add_run(); r.text = text; set_run(r, sz, col, bd)
    return tb


def picture(slide, path, x, y, w):
    slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


# ══════════════════════════════════════════════
# 발표용 PPT 만들기
# ══════════════════════════════════════════════
PRES = [
    ("cover", None, None, None),
    ("problem", "왜 만들었을까요?", ["키오스크가 늘었어요", "그런데 어떤 분들은 사용이 어려워요",
        "고령층·장애인·어린이·외국인 …"], "우리가 도와주자!"),
    ("solution", "우리의 해결책", ["카메라와 AI가 '나이대'를 알아봐요",
        "화면이 스스로 그 사람에 맞게 변신!", "하드웨어(기계)는 안 바꿔요"], "소프트웨어로 해결!"),
    ("bigpic", "큰 그림 한 장", ["① 카메라로 보고", "② 두뇌(AI)가 생각하고", "③ 화면이 바뀐다"],
        "카메라 → 두뇌 → 화면"),
    ("four", "우리 작품의 3가지 마법", ["나이 맞춤 화면",
        "손으로 비접촉 주문", "말로 다국어 주문"], "이 3가지를 기억!"),
    ("privacy", "카메라는 나이대만 봐요", ["누가 왔는지(누구인지)는 몰라요", "얼굴 사진·얼굴 숫자 저장 X",
        "그래서 개인정보가 안전"], "비밀을 지키는 따뜻한 AI"),
    ("age", "나이를 맞히는 AI", ["AI(CNN)가 얼굴 보고 나이대 추측", "어린이 / 어른 / 어르신",
        "→ 알맞은 화면을 골라요"], "나이는 저장 안 하고 화면 결정만"),
    ("variants", "화면이 변신!", ["어르신 모드: 글씨 1.5배·고대비", "어린이 모드: 큰 버튼·쉬운 말",
        "같은 메뉴도 다르게 보여요"], "누구나 보기 쉽게"),
    ("aim", "손으로 가리키기 (에임)", ["손을 들면 커서가 따라와요", "손바닥 중심으로 움직여 안정적",
        "포즈를 바꾸지 않아도 됨"], "손 = 마법 지팡이"),
    ("pinch", "손으로 집기 (핀치)", ["엄지끝 + 검지끝을 붙이면", "가리키던 메뉴가 담겨요",
        "두 손끝이 잘 보여 안정적"], "주먹보다 정확한 담기"),
    ("dwell", "잠깐 기다리기 (드웰)", ["메뉴 위에 커서를 잠깐 두면", "동그란 링이 차올라요",
        "다 차면 자동으로 선택!"], "손 모양이 어려워도 OK"),
    ("filter", "흔들림 부드럽게 (필터)", ["손은 원래 조금 떨려요", "필터가 매끈하게 만들어요",
        "느릴 땐 세게, 빠를 땐 약하게"], "커서가 안 떨려요"),
    ("vote", "여러 번 보고 정하기 (다수결)", ["한 번 보고 정하면 잘못 볼 수 있어요",
        "여러 장면을 모아 '가장 많은 것'", "그래야 안 튀어요"], "투표로 정확하게"),
    ("voice", "말로 주문하기", ["마이크가 말을 글자로 바꿔요(STT)", "글자에서 메뉴·개수·세트를 뽑아요",
        "인터넷 없어도 규칙으로 OK"], "말 → 글자 → 명령"),
    ("langs", "여러 나라 말", ["한국어 / 영어 / 중국어 감지", "영어로 말하면 화면도 영어로",
        "외국인도 편하게"], "다국어 지원"),
    ("brainbody", "두뇌와 몸 (코드 구조)", ["core = 두뇌(비전·수학·자연어·저장)",
        "ui = 몸(창·카드·커서·테마)", "나눠서 만들면 고치기 쉬워요"], "역할을 나눠요"),
    ("threads", "동시에 여러 일 (스레드)", ["카메라 보기는 무거운 일", "따로 일하게 하면 화면이 안 멈춰요",
        "요리사 여러 명처럼"], "화면이 안 멈추는 비결"),
    ("db", "기억과 비밀 (저장소)", ["작은 서랍장 SQLite 사용", "익명 주문 통계를 저장",
        "얼굴·사진은 저장 안 함"], "가볍고 안전한 저장"),
    ("states", "화면 단계 (상태머신)", ["환영 → 메뉴 → 완료", "정해진 규칙으로 단계 이동",
        "90초 가만히 있으면 처음으로"], "신호등처럼 단계가 바뀜"),
    ("demo", "카메라 없어도 OK", ["카메라·인터넷이 없어도", "데모 모드로 모든 기능 시연",
        "그래서 대회에서도 든든"], "언제나 멈추지 않아요"),
    ("points", "컴퓨터가 보는 '점'", ["얼굴은 468개 점", "손은 21개 점(관절)",
        "이 점들로 계산해요"], "숫자로 자랑! 468 / 21"),
    ("flow_all", "전체 순서도", ["시작 → 얼굴·손 보기 → 화면 변신", "→ 주문(손·말) → 결제 → 완료",
        "이 흐름을 외우면 발표 끝!"], "한 흐름으로 설명하기"),
    ("recap", "핵심 5줄 정리", ["이 5줄이면 우리 작품 설명 끝!"], "친구에게 설명해 보기"),
    ("end", None, None, None),
]

SCENE_FN = {
    "cover": sc_cover, "problem": sc_problem, "solution": sc_solution, "bigpic": sc_bigpic,
    "four": sc_four, "privacy": sc_privacy,
    "age": sc_age, "variants": sc_variants, "aim": sc_aim, "pinch": sc_pinch, "dwell": sc_dwell,
    "filter": sc_filter, "vote": sc_vote, "voice": sc_voice, "langs": sc_langs,
    "brainbody": sc_brainbody, "threads": sc_threads, "db": sc_db, "states": sc_states,
    "demo": sc_demo, "points": sc_points, "flow_all": sc_flow_all, "recap": sc_recap, "end": sc_end,
}


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    for key, title, bullets, keyline in PRES:
        img = SCENE_FN[key]()
        s = blank(prs)
        if title is None:                        # 표지/끝 = 전체 그림
            bg(s, "#DCEBFF")
            picture(s, img, 0, 0, 13.333)
            continue
        bg(s, BG)
        band(s, title, PRIMARY, fs=30)
        picture(s, img, 5.55, 1.45, 7.55)        # 오른쪽 큰 그림
        textbox(s, 0.45, 1.7, 5.0, 4.2,
                [{"t": "• " + b, "s": 19} for b in bullets], anchor=MSO_ANCHOR.TOP)
        from pptx.enum.shapes import MSO_SHAPE
        rib = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), Inches(6.55),
                                 Inches(12.5), Inches(0.72))
        rib.fill.solid(); rib.fill.fore_color.rgb = RGB("#FFF1E6"); rib.line.color.rgb = RGB(ACCENT)
        tf = rib.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = "🔑 핵심: " + keyline; set_run(r, 18, "#8a5a12", True)
    prs.save("발표용.pptx")
    print("SAVED 발표용.pptx  slides:", len(prs.slides._sldIdLst))


# ══════════════════════════════════════════════
# 퀴즈 100 (10주제 × 10)
# ══════════════════════════════════════════════
TOPICS = [
    ("배리어프리 모드", PRIMARY), ("카메라·개인정보", GREEN), ("나이 맞히기", ACCENT),
    ("손동작", PINK), ("흔들림 잡기", "#7A5AF8"), ("말로 주문", "#00A5B5"),
    ("프로그램 구조", "#E0455E"), ("동시에 일하기", "#2D6CDF"),
    ("기억과 비밀", "#C08A00"), ("전체 흐름", "#2BB673"),
]

# (질문, 보기 or None, 정답, 한줄해설)
QUIZ = {
 0: [("모두가 쉽게 쓰도록 만든 것을 무엇이라 해요?", ["배리어프리", "게임", "숙제"], "①", "장벽(barrier)이 없다는 뜻."),
     ("어르신에게는 글씨가 커지는 ○○ 모드가 있다.", None, "실버", "글자 1.5배."),
     ("실버 모드는 글씨가 몇 배 커지나요?", ["1.5배", "10배", "그대로"], "①", "1.5배 확대."),
     ("어린이 모드의 특징은?", ["작은 글씨", "큰 버튼·쉬운 말", "어두운 화면"], "②", "쉽게 쓰도록."),
     ("눈이 잘 안 보이는 분을 위한 모드는?", ["고대비 모드", "게임 모드", "비행 모드"], "①", "색을 또렷하게."),
     ("화면 모드는 사람에 맞게 스스로 바뀐다.", None, "O", "자동 변신."),
     ("모드를 손으로 직접 고를 수도 있다.", None, "O", "버튼으로 선택 가능."),
     ("기본 화면(보통 사람용)을 ○○ 모드라고 한다.", None, "일반", "standard 모드."),
     ("배리어프리는 특정한 사람만 위한 것이다.", None, "X", "모두를 위한 것."),
     ("화면이 바뀌어도 파는 메뉴는 똑같다.", None, "O", "보이는 방식만 달라져요.")],
 1: [("카메라는 '누가 왔는지(이름)'를 알아낸다.", None, "X", "나이대만 봐요."),
     ("카메라가 보는 것은?", ["나이대", "이름", "전화번호"], "①", "어린이·어른·어르신."),
     ("우리 키오스크는 얼굴 사진을 저장한다.", None, "X", "아무것도 저장 안 해요."),
     ("나이 정보는 저장하나요?", ["아니요", "네 평생", "네 이름과 함께"], "①", "화면 결정에만 쓰고 안 저장."),
     ("카메라 없이도 프로그램이 동작한다.", None, "O", "데모 모드."),
     ("개인정보를 적게 모으는 원칙을 무엇이라 해요?", ["최소 수집", "많이 모으기", "자랑하기"], "①", "필요한 것만."),
     ("카메라는 특정한 사람을 콕 집어 찾아낸다.", None, "X", "개인 식별은 안 해요."),
     ("얼굴로 '단골'을 자동으로 찾는 기능이 지금 있다.", None, "X", "정확도·개인정보 문제로 뺐어요."),
     ("나이대는 한 번 보고 바로 정한다.", None, "X", "여러 번 보고 다수결."),
     ("카메라 정보로 하는 일은?", ["화면을 알맞게 바꾸기", "돈 받기", "문 열기"], "①", "테마 결정만.")],
 2: [("나이대를 무엇이 맞히나요?", ["AI(CNN)", "마법", "점쟁이"], "①", "학습된 신경망."),
     ("나이를 맞히면 화면이 어린이/어른/어르신용으로 변신한다.", None, "O", "자동 변신."),
     ("어르신 모드는 글씨가 몇 배 커지나?", None, "1.5", "1.5배 확대."),
     ("나이를 매 순간 다시 계산해 화면이 자주 깜빡인다.", None, "X", "다수결·쿨다운으로 안정."),
     ("나이 AI는 몇 그룹으로 나누나요?", ["2", "3", "10"], "②", "어린이·어른·어르신."),
     ("나이는 저장하지 않고 화면 결정에만 쓴다.", None, "O", "개인정보 보호."),
     ("사진을 AI가 보기 좋게 만든 입력을?", ["블롭(blob)", "물방울", "사탕"], "①", "227×227 blob."),
     ("어두울 때 밝기를 고르게 펴는 것을 밝기 ○○이라 한다.", None, "보정", "히스토그램 평활화."),
     ("나이 AI가 없어도 앱은 안 죽고 대체 방법으로 동작한다.", None, "O", "견고성 설계."),
     ("어린이 모드의 특징은?", ["작은 글씨", "큰 버튼·쉬운 말", "어두운 화면"], "②", "쉽게 쓰도록.")],
 3: [("커서를 손으로 움직이는 것을?", ["에임", "점프", "클릭"], "①", "가리키기."),
     ("엄지와 검지를 붙여 담는 것을?", ["핀치", "박수", "주먹"], "①", "집기."),
     ("메뉴 위에 잠깐 머물러 선택하는 것을?", ["드웰", "스크롤", "흔들기"], "①", "머무르기."),
     ("드웰은 동그란 게이지가 다 차면 선택된다.", None, "O", "진행 링."),
     ("손바닥을 좌우로 휘두르면 카테고리가 넘어간다.", None, "O", "스와이프."),
     ("커서는 무엇을 따라 움직여 더 안정적일까?", ["손끝 하나", "손바닥 중심", "팔꿈치"], "②", "여러 점 평균."),
     ("주먹은 손가락이 가려져 잘 안 보일 때가 있다.", None, "O", "occlusion 문제."),
     ("핀치는 두 손끝이 잘 보여 주먹보다 안정적이다.", None, "O", "거리 기반."),
     ("손 관절 점은 모두 몇 개?", None, "21", "MediaPipe 손 21점."),
     ("담는 방법 두 가지는?", ["핀치·드웰", "박수·점프", "소리·빛"], "①", "집기·머무르기.")],
 4: [("손 떨림을 부드럽게 하는 것을?", ["필터", "지우개", "자석"], "①", "잡음 제거."),
     ("여러 번 보고 가장 많이 나온 것으로 정하는 것을 다수결이라 한다.", None, "O", "투표."),
     ("문턱(기준)을 하나만 두면 경계에서 자꾸 깜빡인다.", None, "O", "그래서 두 개 사용."),
     ("켤 때와 끌 때 기준을 다르게 두는 방법은?", ["히스테리시스", "도미노", "미로"], "①", "이중 임계값."),
     ("커서를 부드럽게 하는 특별한 필터 이름은 ○○ 필터.", None, "원-유로", "One-Euro."),
     ("필터를 너무 세게 걸면 반응이 느려질 수 있다.", None, "O", "균형이 중요."),
     ("다수결은 어떤 문제를 막나요?", ["순간 잘못 인식", "배터리", "소리"], "①", "튐 방지."),
     ("손이 멀든 가깝든 똑같이 되도록 손 크기로 기준을 맞춘다.", None, "O", "거리 정규화."),
     ("같은 동작을 너무 자주 실행하지 않게 두는 시간을?", ["쿨다운", "알람", "낮잠"], "①", "연타 방지."),
     ("흔들림을 잡으면 메뉴를 더 정확히 고를 수 있다.", None, "O", "정확도 향상.")],
 5: [("말을 글자로 바꾸는 것을?", ["음성인식(STT)", "TTS", "GPS"], "①", "Speech To Text."),
     ("글자를 소리로 읽어 주는 것을?", ["TTS", "STT", "USB"], "①", "Text To Speech."),
     ("인터넷이 없어도 규칙 방법으로 주문을 알아들을 수 있다.", None, "O", "오프라인 규칙 파서."),
     ("한글이 있으면 언어를 ○○○로 감지한다.", None, "한국어", "ko로 감지."),
     ("영어로 말하면 화면도 영어로 바뀔 수 있다.", None, "O", "다국어 전환."),
     ("복잡한 문장에서 무엇을 뽑아내나요?", ["개수·세트·옵션", "날씨", "색깔"], "①", "주문 정보."),
     ("음성 안내는 눈이 불편한 분께 도움이 된다.", None, "O", "접근성."),
     ("자연어 주문을 도와주는 똑똑한 AI는?", ["GPT(있으면)", "계산기", "시계"], "①", "LLM."),
     ("마이크가 없으면 예시 문장을 골라 담을 수도 있다.", None, "O", "제스처로 선택."),
     ("세 가지 언어는 한국어, 영어, ○○○.", None, "중국어", "ko/en/zh.")],
 6: [("화면(몸)을 담당하는 폴더는?", ["ui", "core", "data"], "①", "프론트엔드."),
     ("두뇌(로직)를 담당하는 폴더는?", ["core", "ui", "tests"], "①", "백엔드."),
     ("두뇌와 몸을 나누면 고치기 쉽다.", None, "O", "모듈화."),
     ("다른 일꾼끼리 소식을 전하는 방법은?", ["시그널/슬롯", "이메일", "전화번호"], "①", "Qt 신호."),
     ("프로그램을 시작하는 파일은 ○○○○.py.", None, "main", "진입점."),
     ("설정·숫자들은 config.py 한 곳에 모여 있다.", None, "O", "쉽게 조정."),
     ("화면 부품(카드·커서)이 있는 파일은?", ["widgets.py", "vision.py", "nlu.py"], "①", "UI 부품."),
     ("기능마다 클래스로 나눈 것을 모듈화라 한다.", None, "O", "OOP."),
     ("얼굴·손을 분석하는 파일은?", ["vision.py", "theme.py", "order.py"], "①", "비전 엔진."),
     ("파일 구조(트리)를 보면 무엇이 어디 있는지 알기 쉽다.", None, "O", "구조 파악.")],
 7: [("여러 일을 동시에 하도록 나눈 것을?", ["스레드", "리본", "사슬"], "①", "동시 실행."),
     ("카메라 처리를 따로 하면 화면이 안 멈춘다.", None, "O", "UI 반응 유지."),
     ("카메라를 담당하는 스레드 이름은?", ["VisionThread", "FastThread", "MainThread"], "①", "비전 스레드."),
     ("다른 스레드가 화면을 직접 만지면 문제가 생길 수 있다.", None, "O", "그래서 신호로 전달."),
     ("그래서 결과만 ○○○로 안전하게 전달한다.", None, "시그널", "signal/slot."),
     ("요리사 여러 명이 동시에 요리하는 것과 비슷하다.", None, "O", "좋은 비유."),
     ("파이썬에서 한 번에 한 스레드만 코드를 도는 규칙은?", ["GIL", "HTML", "RGB"], "①", "전역 인터프리터 락."),
     ("그래도 카메라·영상처리는 GIL을 풀어 동시에 도움이 된다.", None, "O", "C 확장·I/O."),
     ("신호는 순서대로 처리되어 안전하다.", None, "O", "큐 처리."),
     ("스레드를 나누는 가장 큰 이유는?", ["화면이 안 멈추게", "소리 크게", "색깔"], "①", "부드러운 화면.")],
 8: [("주문 기록을 기억하는 작은 저장소는?", ["SQLite", "유튜브", "카메라"], "①", "파일 하나 DB."),
     ("저장소에 얼굴 사진을 저장한다.", None, "X", "얼굴은 저장 안 해요."),
     ("저장소에는 익명 ○○ 통계만 저장한다.", None, "주문", "익명 저장."),
     ("데이터는 매장 안(로컬)에 저장되어 밖으로 안 나간다.", None, "O", "프라이버시."),
     ("주문 통계로 무엇을 알 수 있나요?", ["요일·시간대 인기", "날씨", "키"], "①", "트렌드 분석."),
     ("두 일꾼이 같은 저장소를 쓸 때 잠금(Lock)으로 안전하게 한다.", None, "O", "동시성 보호."),
     ("SQLite의 장점이 아닌 것은?", ["설치 필요 없음", "파일 하나", "수백만 명 웹 처리"], "③", "대규모 웹용 아님."),
     ("개인정보를 지키는 것은 '따뜻한 AI'의 중요한 약속이다.", None, "O", "핵심 가치."),
     ("카메라 영상은 밖으로 보내지 않고 안에서만 처리한다.", None, "O", "오프라인 처리."),
     ("얼굴 정보를 저장하지 않아 개인정보가 안전하다.", None, "O", "얼굴 미저장.")],
 9: [("화면 단계 순서는?", ["환영→메뉴→완료", "완료→메뉴", "메뉴만"], "①", "상태머신."),
     ("화면 단계가 정해진 규칙으로 바뀌는 것을 상태머신이라 한다.", None, "O", "FSM."),
     ("카메라가 없으면 앱이 꺼진다.", None, "X", "데모 모드로 동작."),
     ("아무 동작이 없으면 몇 초 후 첫 화면으로 돌아가나?", None, "90", "세션 타임아웃."),
     ("결제하면 무엇을 하나요?", ["주문번호 발급·저장", "전원 끔", "삭제"], "①", "DB에 기록."),
     ("데모 모드에서는 버튼·음성으로 모든 기능을 보여줄 수 있다.", None, "O", "항상 시연 가능."),
     ("얼굴 점 468개, 손 점 몇 개?", None, "21", "손 21점."),
     ("우리 작품의 핵심 가치가 아닌 것은?", ["자동 변신", "개인정보 보호", "광고 많이"], "③", "광고 아님."),
     ("인터넷이 끊겨도 규칙 방법으로 주문을 계속할 수 있다.", None, "O", "오프라인 견고성."),
     ("이 모든 걸 한마디로?", ["누구나 쉬운 배리어프리 키오스크", "게임기", "냉장고"], "①", "핵심 요약.")],
}


def q_lines(items, start):
    """문제 5개를 슬라이드 텍스트 항목으로."""
    out = []
    for k, (q, opts, ans, why) in enumerate(items):
        num = start + k
        typ = "(O/X)" if opts is None and ans in ("O", "X") else ("(빈칸)" if opts is None else "")
        out.append({"t": f"{num}. {q} {typ}".rstrip(), "s": 18, "c": INK, "b": True})
        if opts:
            line = "   " + "   ".join(f"{'①②③④'[j]} {o}" for j, o in enumerate(opts))
            out.append({"t": line, "s": 16, "c": "#3C4a5c", "b": False})
    return out


def a_lines(items, start):
    out = []
    for k, (q, opts, ans, why) in enumerate(items):
        num = start + k
        out.append({"t": f"{num}. ✅ {ans}  —  {why}", "s": 17, "c": GREEN, "b": True})
    return out


def build_quiz():
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    # 표지
    s = blank(prs); bg(s, "#DCEBFF")
    picture(s, sc_cover(), 0, 0, 13.333)
    tb = s.shapes.add_textbox(Inches(0), Inches(0.3), Inches(13.333), Inches(1.2))
    p = tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = "🎯 퀴즈 100문제"; set_run(r, 34, PRIMARY, True)
    # 사용법
    s = blank(prs); bg(s, BG); band(s, "퀴즈 사용법", ACCENT, fs=30)
    textbox(s, 1.0, 2.0, 11.3, 4.5, [
        {"t": "• 10개 주제 × 10문제 = 총 100문제", "s": 22},
        {"t": "• 문제를 먼저 풀고, 바로 다음 슬라이드에서 정답을 확인해요.", "s": 22},
        {"t": "• 유형: O/X · 3지선다(①②③) · 빈칸 채우기", "s": 22},
        {"t": "• 3명이 각자 풀고 서로 채점해 보세요!", "s": 22},
    ], anchor=MSO_ANCHOR.MIDDLE)
    msc = mascot_small()
    # 주제별
    for ti, (title, color) in enumerate(TOPICS):
        banner = topic_banner(ti + 1, title, color)
        # 주제 구분 슬라이드
        s = blank(prs); bg(s, "#FFFFFF"); picture(s, banner, 0, 2.7, 13.333)
        items = QUIZ[ti]
        for half in (0, 5):
            grp = items[half:half + 5]
            # 문제
            s = blank(prs); bg(s, BG)
            band(s, f"[{title}] 문제 {half+1}~{half+5}", color, fs=26)
            textbox(s, 0.7, 1.5, 10.6, 5.6, q_lines(grp, half + 1), anchor=MSO_ANCHOR.TOP)
            picture(s, msc, 11.5, 5.45, 1.5)
            # 정답
            s = blank(prs); bg(s, "#F2FBF5")
            band(s, f"[{title}] 정답 {half+1}~{half+5}", GREEN, fs=26)
            textbox(s, 0.9, 1.7, 10.4, 5.2, a_lines(grp, half + 1), anchor=MSO_ANCHOR.TOP)
            picture(s, msc, 11.5, 5.45, 1.5)
    prs.save("퀴즈100.pptx")
    print("SAVED 퀴즈100.pptx  slides:", len(prs.slides._sldIdLst))


def main():
    if os.path.isdir(ASSET):
        shutil.rmtree(ASSET)
    os.makedirs(ASSET, exist_ok=True)
    build_presentation()
    build_quiz()


if __name__ == "__main__":
    main()
